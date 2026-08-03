"""The single entry point: ``prepare(bytes, mime, filename) -> PreparedContent``.

A pure orchestrator — format dispatch + parse/decode + hybrid-OCR routing +
provenance normalization + truncation. No storage, no chunking, no embedding, no
async, no business meaning. The only I/O is the injected OCR provider.

Parsing mirrors the ``rakam_systems`` vectorstore loaders (``pymupdf4llm`` for
PDFs, ``openpyxl`` for spreadsheets) without importing that heavy package, so
this stays a light dependency.
"""
from __future__ import annotations

import base64  # noqa: F401  (kept for symmetry with providers; harmless)
import csv
import io

from .providers import OCRProvider
from .schema import (
    PAGE_DELIMITER_RE,
    PreparedContent,
    SourceRef,
    TableRow,
    page_delimiter,
)

# Below this many stripped characters a PDF page is treated as "no text layer"
# (i.e. scanned) and routed to OCR.
_PDF_TEXT_MIN = 20


def _join_pages(pages: list[str]) -> str:
    """Delimit per-page markdown so a consumer can slice it back apart."""
    return "\n\n".join(
        f"{page_delimiter(i)}\n{text}" for i, text in enumerate(pages, start=1)
    )


def _truncate(markdown: str, max_chars: int) -> tuple[str, bool, int | None]:
    """Cut to ``max_chars`` without severing a page delimiter.

    A naive slice can land inside ``<!-- page:12 -->`` and leave a fragment
    that breaks the consumer's regex, so prefer the last *whole* page that
    fits. When even the first page overruns, cut mid-page — but never inside a
    marker. Returns ``(markdown, truncated, last_whole_page)``."""
    if len(markdown) <= max_chars:
        return markdown, False, None

    last_boundary = None
    last_page = None
    for match in PAGE_DELIMITER_RE.finditer(markdown):
        if match.start() > max_chars:
            break
        last_boundary, last_page = match.start(), int(match.group(1))

    # A boundary at 0 is the first page's own marker — cutting there would
    # yield an empty document, so fall through to the mid-page cut.
    if last_boundary:
        return markdown[:last_boundary].rstrip("\n"), True, last_page - 1

    cut = markdown[:max_chars]
    # Drop a delimiter the slice may have bisected.
    partial = cut.rfind("<!-- page:")
    if partial != -1 and "-->" not in cut[partial:]:
        cut = cut[:partial].rstrip("\n")
    return cut, True, None


def prepare(
    content: bytes,
    mime: str,
    filename: str,
    ocr: OCRProvider | None = None,
    max_chars: int = 8000,
) -> PreparedContent:
    """Turn one document into ``PreparedContent``. Never raises — an unreadable
    file yields ``provider="none"`` + a ``meta['help']`` hint."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if mime == "application/pdf" or ext == "pdf" or content[:5] == b"%PDF-":
            pc = _prepare_pdf(content, mime, ocr)
        elif mime.startswith("image/") or ext in ("png", "jpg", "jpeg", "tiff", "webp"):
            pc = _prepare_image(content, mime, ocr)
        elif ext in ("xlsx", "xlsm") or "spreadsheetml" in mime or "ms-excel" in mime:
            pc = _prepare_xlsx(content)
        elif ext == "csv" or mime == "text/csv":
            pc = _prepare_csv(content)
        elif ext == "eml" or mime == "message/rfc822":
            pc = _prepare_email(content)
        else:
            text = _decode_text(content)
            pc = PreparedContent(
                markdown=text or "",
                provider="text" if text else "none",
                meta={} if text else {"help": "could not read this file (binary?)"},
            )
    except Exception as exc:  # noqa: BLE001 — best-effort: a bad file must never raise
        pc = PreparedContent(provider="none", meta={"help": f"processing failed: {exc}"})

    # Recorded before the cut so a consumer can say "showing 8000 of 143000"
    # instead of lying by omission.
    pc.meta["total_chars"] = len(pc.markdown)
    pc.markdown, pc.truncated, last_whole_page = _truncate(pc.markdown, max_chars)
    if last_whole_page is not None:
        pc.meta["truncated_after_page"] = last_whole_page
    return pc


# ── PDF (hybrid: native text first, OCR for scans) ─────────────────────────
def _pdf_text_len(content: bytes) -> tuple[int, int]:
    import fitz

    with fitz.open(stream=content, filetype="pdf") as doc:
        total = sum(len(page.get_text().strip()) for page in doc)
        return total, doc.page_count


def _strip_page_rule(text: str) -> str:
    """Drop pymupdf4llm's trailing ``-----`` page separator from a chunk."""
    return text.rstrip().removesuffix("-----").rstrip()


# A table-heavy PDF (a price list, a long packing list) can carry thousands of
# rows, and ``rows`` is persisted by consumers. Cap it and say so in ``meta``
# rather than silently returning a subset.
_MAX_TABLE_ROWS = 2000


def _column_names(names: list, width: int) -> list[str]:
    """Header labels for a table, one per column.

    Real documents give blank, ``None`` and duplicate headers — a supplier
    acknowledgement routinely has unnamed spacer columns. Blanks become
    ``colN`` (matching the spreadsheet/CSV branches) and collisions are
    suffixed, so ``cells`` never silently loses a column to a dict-key clash."""
    out: list[str] = []
    seen: dict[str, int] = {}
    for i in range(width):
        raw = (names[i] if i < len(names) else None) or ""
        # Collapse the embedded newlines pymupdf leaves in wrapped headers.
        name = " ".join(str(raw).split()) or f"col{i}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        out.append(name)
    return out


def _find_tables(page):
    """Detect a page's tables, preferring the strategy that keeps numeric
    columns apart.

    Measured on real supplier documents (an order acknowledgement and a packing
    list):

    * ``lines_strict`` splits a line item into its own columns —
      ``['2 Bloc-Porte ...', '3 U', '327.78 €', '983.34 €']`` — which is what
      comparing a quantity against an ERP quantity requires.
    * ``lines`` finds the same table but collapses quantity, unit price and
      total into the description cell, leaving the numeric columns empty.
    * ``text`` shatters a row into ~10 fragment columns; unusable.

    ``lines_strict`` mangles the *header* (a wrapped multi-line header lands in
    column 0), so callers fall back to positional ``colN`` names — the semantic
    mapping from column to business field is client-specific anyway.
    """
    tables = page.find_tables(strategy="lines_strict").tables
    if tables:
        return tables
    # A borderless table has no strict ruling to find; the looser strategy at
    # least recovers the rows.
    return page.find_tables(strategy="lines").tables


def _pdf_table_rows(doc) -> tuple[list[TableRow], bool]:
    """Structured line items from a native PDF's tables, with page provenance.

    The markdown already renders these tables for a reader; this is the same
    data as *fields*, which is what line-by-line comparison against an ERP
    needs (quantity vs quantity, price vs price). ``pymupdf4llm``'s own
    ``tables`` key carries only geometry, so the cells come from PyMuPDF's
    ``find_tables()``.

    Best-effort by design: table detection is heuristic, so a page that fails
    to parse is skipped rather than failing the document. Returns
    ``(rows, capped)``."""
    rows_out: list[TableRow] = []
    for pno in range(doc.page_count):
        try:
            tables = _find_tables(doc[pno])
        except Exception:  # noqa: BLE001, PERF203 — heuristic; a bad page must not sink the doc
            continue
        for table in tables:
            try:
                data = table.extract()
            except Exception:  # noqa: BLE001, PERF203
                continue
            if not data:
                continue
            header = getattr(table, "header", None)
            names = list(header.names) if header and header.names else []
            # external=False means the header row is ALSO data[0]; emitting it
            # would produce a row whose values are its own column names.
            body = data[1:] if (header and not header.external) else data
            width = max((len(r) for r in data), default=0)
            columns = _column_names(names, width)
            for r_idx, raw_row in enumerate(body, start=1):
                cells = {
                    columns[i]: " ".join(str(raw_row[i] or "").split())
                    for i in range(min(len(raw_row), width))
                }
                if not any(cells.values()):  # spacer/rule rows
                    continue
                rows_out.append(
                    TableRow(cells=cells, source=SourceRef(page=pno + 1, row=r_idx)),
                )
                if len(rows_out) >= _MAX_TABLE_ROWS:
                    return rows_out, True
    return rows_out, False


def _prepare_pdf(content: bytes, mime: str, ocr: OCRProvider | None) -> PreparedContent:
    text_len, pages = _pdf_text_len(content)
    if text_len >= _PDF_TEXT_MIN:
        import fitz
        import pymupdf4llm

        with fitz.open(stream=content, filetype="pdf") as doc:
            # page_chunks: per-page markdown instead of one fused string, so
            # provenance points at locatable regions rather than being a bare
            # list of page numbers.
            # show_progress: the default prints "Processing ..." and an ASCII
            # progress bar to stdout — on a server that is one log spam burst
            # per uploaded file.
            chunks = pymupdf4llm.to_markdown(
                doc, page_chunks=True, show_progress=False,
            )
            table_rows, capped = _pdf_table_rows(doc)
        # Read the page number from chunk metadata rather than enumerating —
        # do not assume chunk order matches page order for every document.
        numbered = sorted(
            (
                (int((c.get("metadata") or {}).get("page", i + 1)), c.get("text") or "")
                for i, c in enumerate(chunks)
            ),
            key=lambda pair: pair[0],
        )
        return PreparedContent(
            # Strip pymupdf4llm's own trailing "-----" page rule: it would sit
            # alongside our delimiter as a second, ambiguous marker.
            markdown=_join_pages([_strip_page_rule(text) for _, text in numbered]),
            rows=table_rows,
            provider="pdf_text",
            provenance=[SourceRef(page=n) for n, _ in numbered],
            meta={
                "page_count": pages,
                **({"rows_capped": _MAX_TABLE_ROWS} if capped else {}),
            },
        )
    if ocr and ocr.available():
        markdown, provenance = ocr.ocr(content, mime)
        return PreparedContent(
            markdown=markdown,
            provider=f"{ocr.name}_ocr",
            provenance=provenance,
            meta={"page_count": len(provenance) or pages},
        )
    return PreparedContent(
        provider="none",
        meta={"page_count": pages, "help": "no text layer and OCR unavailable"},
    )


def _prepare_image(content: bytes, mime: str, ocr: OCRProvider | None) -> PreparedContent:
    if ocr and ocr.available():
        markdown, provenance = ocr.ocr(content, mime)
        return PreparedContent(
            markdown=markdown,
            provider=f"{ocr.name}_ocr",
            provenance=provenance,
            # Was the one paged branch reporting no page_count at all.
            meta={"page_count": len(provenance) or 1},
        )
    return PreparedContent(provider="none", meta={"help": "image needs OCR; provider unavailable"})


# ── tabular (markdown table + structured rows + provenance) ────────────────
def _rows_to_markdown(headers: list[str], data_rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in data_rows]
    return "\n".join(lines)


def _prepare_xlsx(content: bytes) -> PreparedContent:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet_names = [ws.title for ws in wb.worksheets]
    md_blocks: list[str] = []
    rows_out: list[TableRow] = []
    for ws in wb.worksheets:
        headers: list[str] | None = None
        data_rows: list[list[str]] = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = ["" if c is None else str(c) for c in row]
            if headers is None:
                headers = vals
                continue
            data_rows.append(vals)
            cells = {(headers[i] if i < len(headers) else f"col{i}"): vals[i] for i in range(len(vals))}
            rows_out.append(TableRow(cells=cells, source=SourceRef(sheet=ws.title, row=r_idx)))
        if headers:
            md_blocks.append(f"### {ws.title}\n" + _rows_to_markdown(headers, data_rows))
    wb.close()
    return PreparedContent(
        markdown="\n\n".join(md_blocks),
        rows=rows_out,
        provider="tabular",
        meta={"sheet_names": sheet_names},
    )


def _prepare_csv(content: bytes) -> PreparedContent:
    text = _decode_text(content) or ""
    reader = list(csv.reader(io.StringIO(text)))
    if not reader:
        return PreparedContent(provider="text")
    headers, data = reader[0], reader[1:]
    rows_out = [
        TableRow(
            cells={(headers[i] if i < len(headers) else f"col{i}"): vals[i] for i in range(len(vals))},
            source=SourceRef(row=r_idx),
        )
        for r_idx, vals in enumerate(data, start=2)
    ]
    return PreparedContent(
        markdown=_rows_to_markdown(headers, data),
        rows=rows_out,
        provider="tabular",
        meta={"row_count": len(rows_out)},
    )


# ── email / text ───────────────────────────────────────────────────────────
def _decode_text(content: bytes) -> str | None:
    if b"\x00" in content[:4096]:  # reject binary masquerading as text
        return None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return None


def _prepare_email(content: bytes) -> PreparedContent:
    import email
    from email import policy

    msg = email.message_from_bytes(content, policy=policy.default)
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    subject = msg.get("subject", "")
    markdown = f"# {subject}\n\n**From:** {msg.get('from', '')}\n\n{text}".strip()
    return PreparedContent(markdown=markdown, provider="text", meta={"subject": subject})
